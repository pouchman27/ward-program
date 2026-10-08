"""Offline safety tests. Never writes to live Sheets."""
import unittest,tempfile,os,datetime
import openpyxl
from clean_schema import SAC_COLUMNS,ACT_COLUMNS
from sync_middleware import sac_rows,act_rows,sac_row
from sync_job import merge,cells
from mw_reader import table_rows
class Safety(unittest.TestCase):
 def test_freeze_missing_and_update(self):
  old=[{'Date':'2026-09-01','value':'frozen'},{'Date':'2026-10-04','value':'keep'}]
  new=[{'Date':'2026-09-01','value':'changed'},{'Date':'2026-10-11','value':'new'}]
  self.assertEqual(merge(old,new,'2026-09-23'),old+[new[1]])
 def test_bad_program_tab_retains_row(self):
  wb=openpyxl.Workbook();ws=wb.active;ws.title='Damaged weekly tab';ws['A1']=datetime.datetime(2026,10,11);ws['A3']='Not a recognized program layout'
  with tempfile.TemporaryDirectory() as d:
   path=os.path.join(d,'bad.xlsx');wb.save(path)
   rows,problems,skipped=sac_rows(path,'2026-09-23')
  self.assertEqual(rows,{})
  self.assertTrue(any('layout mismatch' in msg for _,msg in skipped))
  previous=[dict(zip(SAC_COLUMNS,['2026-10-11']+['']*(len(SAC_COLUMNS)-1)))]
  self.assertEqual(merge(previous,list(rows.values()),'2026-09-23'),previous)
 def test_bad_activity_layout_fails_before_write(self):
  wb=openpyxl.Workbook();wb.active.title='2026 Activities';wb.active['A1']='Date';wb.active['B1']='Wrong layout'
  with tempfile.TemporaryDirectory() as d:
   path=os.path.join(d,'bad.xlsx');wb.save(path)
   with self.assertRaisesRegex(ValueError,'header row'):act_rows(path)
 def test_reader_rejects_extra_columns(self):
  with self.assertRaisesRegex(ValueError,'header mismatch'):table_rows([SAC_COLUMNS+['Private']],SAC_COLUMNS)
 def test_reader_rejects_duplicate_dates(self):
  r=['2026-10-11']+['']*(len(SAC_COLUMNS)-1)
  with self.assertRaisesRegex(ValueError,'Duplicate date'):table_rows([SAC_COLUMNS,r,r],SAC_COLUMNS)
if __name__=='__main__':unittest.main()
