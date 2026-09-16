#!/usr/bin/env python3
"""Generate the Milton Ward virtual program site from the sacrament agendas workbook."""
import openpyxl, json, re, datetime, sys, os, glob

XLSX = sys.argv[1] if len(sys.argv) > 1 else '/tmp/agendas.xlsx'
OUT = sys.argv[2] if len(sys.argv) > 2 else '/home/sandbox/wardsite/site/data.json'
SITE_PASSPHRASE = os.environ.get('WARD_SITE_PASSPHRASE', 'MiltonWard')
try:
    HYMN_URLS = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'hymn_map.json')))
except Exception:
    HYMN_URLS = {}
SHEET_URL = 'https://docs.google.com/spreadsheets/d/1fvw_Kyhy44Er_al2Ammw07KbXA39iBporF3bOF_B1aw/edit'

SKIP_TABS = {'Template', 'Sheet20', '2026 Advancements'}
NUM_RE = re.compile(r'^\s*(\d+)\s*:\s*$')
QUOTEY = ('we propose', 'manifest it', 'stand as your name', "bishop's office",
          'released the following', 'have been called', 'vote of thanks')

def s(v):
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): v = int(v)
    if isinstance(v, (datetime.datetime, datetime.date)): return v.strftime('%-m/%-d/%Y')
    return str(v).strip()

def is_noise(t):
    tl = t.lower().strip('"').strip()
    return not tl or any(q in tl for q in QUOTEY) or tl.startswith('"')

class Grid:
    def __init__(self, ws):
        self.ws = ws
        self.cells = {}   # (row,col) -> text
        self.links = {}   # (row,col) -> hyperlink target
        for row in ws.iter_rows():
            for c in row:
                t = s(c.value)
                if t:
                    self.cells[(c.row, c.column)] = t
                    if c.hyperlink and c.hyperlink.target:
                        self.links[(c.row, c.column)] = c.hyperlink.target
        self.max_row = max((r for r, _ in self.cells), default=0)

    def find(self, pattern, flags=re.I):
        rx = re.compile(pattern, flags)
        for (r, c), t in sorted(self.cells.items()):
            if rx.search(t):
                return r, c, t
        return None

    def right_of(self, row, col, max_col=14):
        for cc in range(col + 1, max_col + 1):
            t = self.cells.get((row, cc))
            if t: return t
        return ''

    def region_text(self, r1, r2, drop_markers=True, skip_cells=()):
        """Non-empty lines in row range, reading order, skipping noise/markers."""
        out = []
        for r in range(r1, r2 + 1):
            parts = []
            for c in range(1, 15):
                if (r, c) in skip_cells: continue
                t = self.cells.get((r, c), '')
                if not t or len(t) < 2: continue
                if drop_markers and NUM_RE.match(t): continue
                if is_noise(t): continue
                parts.append(t)
            if parts: out.append(' '.join(parts))
        return out

def parse_hymn(g, row):
    num, title = '', ''
    for c in range(1, 14):
        t = g.cells.get((row, c), '')
        if t.strip() == '#':
            v = g.right_of(row, c)
            if not re.match(r'^name\s*:', v, re.I):
                num = v
        elif re.match(r'^name\s*:', t, re.I):
            title = g.right_of(row, c)
    num = re.sub(r'\.0$', '', num)
    if not (num or title): return None
    out = {'number': num, 'title': title}
    if num and str(num) in HYMN_URLS: out['url'] = HYMN_URLS[str(num)]
    return out

def parse_announcements(g, r_start, r_end):
    markers = []
    for (r, c), t in g.cells.items():
        if not (r_start <= r <= r_end): continue
        m = NUM_RE.match(t)
        if m:
            markers.append((r, c, int(m.group(1))))
        else:
            m2 = re.search(r'(\d+)\s*:\s*$', t)
            if m2 and re.match(r'^(announcements|post)', t, re.I):
                markers.append((r, c, int(m2.group(1))))
    if not markers: return []
    right_cols = [c for _, c, _ in markers if c >= 5]
    boundary = min(right_cols) if right_cols else 14
    items = []
    for grp in ('L', 'R'):
        gms = sorted([m for m in markers if (m[1] < 5) == (grp == 'L')])
        for i, (r, c, n) in enumerate(gms):
            r2 = (gms[i + 1][0] - 1) if i + 1 < len(gms) else r_end
            c1, c2 = c + 1, (boundary - 1 if grp == 'L' else 14)
            lines = []
            for rr in range(r, r2 + 1):
                parts = []
                for cc in range(c1, c2 + 1):
                    t = g.cells.get((rr, cc), '')
                    if t and not NUM_RE.match(t) and not is_noise(t):
                        parts.append(t)
                if parts: lines.append(' '.join(parts))
            if lines: items.append({'num': n, 'text': '\n'.join(lines)})
    items.sort(key=lambda x: x['num'])
    # dedupe repeated numbers (keep first non-empty)
    seen, out = set(), []
    for it in items:
        if it['num'] in seen: continue
        seen.add(it['num']); out.append(it)
    return out

def parse_business(g, r_start, r_end):
    releases, sustainings = [], []
    for r in range(r_start, r_end + 1):
        # releases: name col A, 'From' col C, calling col D
        name = g.cells.get((r, 1), '')
        if name and not is_noise(name) and g.cells.get((r, 3), '').lower() == 'from':
            calling = g.cells.get((r, 4), '')
            releases.append({'name': name, 'calling': calling})
        # sustainings: name col G/H, 'To' then calling
        nm = ''
        nc = 0
        for c in (7, 8):
            t = g.cells.get((r, c), '')
            if t and not is_noise(t) and t.lower() != 'to':
                nm, nc = t, c; break
        if nm:
            calling = ''
            for cc in range(nc + 1, 15):
                t = g.cells.get((r, cc), '')
                if not t: continue
                if t.lower() == 'to' or is_noise(t): continue
                calling = t; break
            sustainings.append({'name': nm, 'calling': calling})
    return releases, sustainings

def parse_advancements(g):
    def col_list(col, r1, r2):
        groups, cur = [], None
        for r in range(r1, r2 + 1):
            t = g.cells.get((r, col), '')
            if not t or is_noise(t): continue
            tl = t.lower()
            if re.search(r'class|teachers|priests|deacons|young (women|men)', tl):
                cur = {'group': t, 'names': []}; groups.append(cur)
            elif cur:
                cur['names'].append(t)
        return groups
    def presidencies(name_col, role_col, r1, r2):
        groups, cur = [], None
        for r in range(r1, r2 + 1):
            t = g.cells.get((r, name_col), '')
            role = g.cells.get((r, role_col), '')
            if not t or is_noise(t): continue
            if not role:
                cur = {'group': t, 'members': []}; groups.append(cur)
            elif cur:
                cur['members'].append({'name': t, 'role': role})
        return groups
    return {
        'young_women': col_list(1, 2, 25),
        'young_men': col_list(3, 2, 14),
        'priesthood_presidencies': presidencies(7, 8, 2, 25),
        'yw_presidencies': presidencies(10, 11, 2, 25),
    }

def parse_program(ws):
    g = Grid(ws)
    d = g.find(r'^\d{4}-\d{2}-\d{2}|^\d{1,2}/\d{1,2}/\d{4}')
    date = None
    for (r, c), t in g.cells.items():
        if r <= 2:
            raw = ws.cell(row=r, column=c).value
            if isinstance(raw, (datetime.datetime, datetime.date)):
                date = raw.date() if isinstance(raw, datetime.datetime) else raw
                break
    if not date: return None
    prog = {'date': date.isoformat(),
            'slug': date.isoformat(),
            'tab': ws.title}
    # bulletin link omitted per Nathan, 2026-09-06
    p = g.find(r'^presiding$')
    if p: prog['presiding'] = g.right_of(p[0], p[1])
    cd = g.find(r'^conducting')
    if cd: prog['conducting'] = g.right_of(cd[0], cd[1])
    ro = g.find(r'recognize others')
    if ro: prog['recognize'] = g.right_of(ro[0], ro[1])
    wv = g.find(r'^welcome visitors')
    if wv: prog['welcome'] = g.right_of(wv[0], wv[1])
    oh = g.find(r'opening hymn')
    if oh: prog['opening_hymn'] = parse_hymn(g, oh[0])
    inv = g.find(r'^invocation')
    if inv: prog['invocation'] = g.right_of(inv[0], inv[1])
    og = g.find(r'thank the organist')
    if og: prog['organist'] = g.right_of(og[0], og[1])
    ch = g.find(r'and chorister')
    if ch: prog['chorister'] = g.right_of(ch[0], ch[1])
    mr = g.find(r'membership records received')
    an = g.find(r'^announcements')
    wb_ = g.find(r'^ward business$')
    st = g.find(r'turn time to stake')
    od = g.find(r'ordinances and recognitions')
    bb = g.find(r'^baby blessing')
    sh = g.find(r'sacrament hymn')
    ad = g.find(r'administration of the sacrament')
    ih = g.find(r'intermediate hymn')
    mn = g.find(r'special musical number')
    cl = g.find(r'closing hymn')
    be = g.find(r'^benediction')
    # membership records, ordinances, baby blessing omitted per Nathan, 2026-09-06
    # Ward business (releases/sustainings) and stake business intentionally
    # omitted from the public site per Nathan, 2026-09-06.
    # Announcements restored for the dedicated tab per Nathan, 2026-09-06.
    if an:
        nxt = [x[0] for x in (wb_, st, od, bb, sh) if x and x[0] > an[0]]
        anns = parse_announcements(g, an[0], (min(nxt) - 1) if nxt else an[0] + 25)
        if anns: prog['announcements'] = anns
    if sh: prog['sacrament_hymn'] = parse_hymn(g, sh[0])
    speakers = []
    for n in (1, 2, 3, 4):
        sp = g.find(rf'^speaker {n}\s*:')
        if sp:
            v = g.right_of(sp[0], sp[1])
            if v: speakers.append({'n': n, 'name': v})
    if speakers: prog['speakers'] = speakers
    if ih: prog['intermediate_hymn'] = parse_hymn(g, ih[0])
    if mn:
        entry = {}
        for rr in range(mn[0], mn[0] + 4):
            for cc in range(1, 14):
                t = g.cells.get((rr, cc), '')
                if re.match(r'^song name', t, re.I): entry['song'] = g.right_of(rr, cc)
                elif re.match(r'^singer name', t, re.I): entry['singers'] = g.right_of(rr, cc)
                elif re.match(r'^accompan', t, re.I): entry['accompanied'] = g.right_of(rr, cc)
        entry = {k: v for k, v in entry.items() if v}
        if entry: prog['musical_number'] = entry
    if cl: prog['closing_hymn'] = parse_hymn(g, cl[0])
    if be: prog['benediction'] = g.right_of(be[0], be[1])
    return prog


ACT_XLSX_URL = 'https://docs.google.com/spreadsheets/d/1_6NKa_YWGWKSLCINSjZeAUrhwWAj4hBrWWF88O7Qj-I/export?format=xlsx'
ACT_SHEET_URL = 'https://docs.google.com/spreadsheets/d/1_6NKa_YWGWKSLCINSjZeAUrhwWAj4hBrWWF88O7Qj-I/edit'
GROUPS = ['YW 12/13', 'YW 14/15', 'YW 16+', 'Deacons', 'Teachers', 'Priests']
MONTHS = {m.lower(): i+1 for i, m in enumerate(['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sept','Oct','Nov','Dec'])}

def parse_activities(path):
    wb = openpyxl.load_workbook(path)
    if '2026 Activities' not in wb.sheetnames: return None
    ws = wb['2026 Activities']
    rows = []
    month = None
    for r in range(9, 64):
        a = s(ws.cell(row=r, column=1).value)
        if a:
            mname = a.split('\n')[0].strip().rstrip('(').strip()
            key = mname.lower()[:4]
            month = MONTHS.get(key) or MONTHS.get(key[:3]) or month
        dayv = ws.cell(row=r, column=2).value
        try:
            day = int(float(dayv))
        except (TypeError, ValueError):
            continue
        if not month: continue
        date = datetime.date(2026, month, day)
        groups = {}
        for gi, gname in enumerate(GROUPS):
            v = s(ws.cell(row=r, column=7 + gi).value)
            if v: groups[gname] = v
        rows.append({
            'date': date.isoformat(),
            'schedule': s(ws.cell(row=r, column=3).value),
            'notes': s(ws.cell(row=r, column=5).value),
            'reference': s(ws.cell(row=r, column=6).value),
            'groups': groups,
        })
    return {'rows': rows, 'sheet_url': ACT_SHEET_URL}


YW_CLASS_NAMES = {'YW 12/13': 'Builders of Faith', 'YW 14/15': 'Messengers of Hope', 'YW 16+': 'Gatherers of Light'}
CAL_LOCATION = '1750 Windward Concourse, Alpharetta, GA'

def ics_escape(t):
    return t.replace('\\', '\\\\').replace(';', '\\;').replace(',', '\\,').replace('\n', '\\n')

def ics_fold(line):
    out = []
    while len(line.encode('utf-8')) > 74:
        cut = 74
        while cut > 0 and (line.encode('utf-8')[:cut+1].decode('utf-8', 'ignore') != line[:cut+1]):
            cut -= 1
        out.append(line[:cut])
        line = ' ' + line[cut:]
    out.append(line)
    return out

def parse_time_min(s):
    m = re.search(r'(\d{1,2})(?::(\d{2}))?\s*(a\.?m\.?|p\.?m\.?)', s or '', re.I)
    if not m: return None
    hh = int(m.group(1)) % 12
    if m.group(3).lower().startswith('p'): hh += 12
    return hh * 60 + int(m.group(2) or 0)

def build_ics(programs, act):
    now = datetime.datetime.now(datetime.timezone.utc).strftime('%Y%m%dT%H%M%SZ')
    lines = ['BEGIN:VCALENDAR', 'VERSION:2.0', 'PRODID:-//Milton Ward//Ward Program//EN',
             'CALSCALE:GREGORIAN', 'METHOD:PUBLISH', 'X-WR-CALNAME:Milton Ward',
             'X-WR-CALDESC:Sacrament meeting programs and youth activities for Milton Ward',
             'BEGIN:VTIMEZONE', 'TZID:America/New_York',
             'BEGIN:DAYLIGHT', 'TZOFFSETFROM:-0500', 'TZOFFSETTO:-0400', 'TZNAME:EDT',
             'DTSTART:19700308T020000', 'RRULE:FREQ=YEARLY;BYMONTH=3;BYDAY=2SU', 'END:DAYLIGHT',
             'BEGIN:STANDARD', 'TZOFFSETFROM:-0400', 'TZOFFSETTO:-0500', 'TZNAME:EST',
             'DTSTART:19701101T020000', 'RRULE:FREQ=YEARLY;BYMONTH=11;BYDAY=1SU', 'END:STANDARD',
             'END:VTIMEZONE']
    def ev(uid, date, start_min, end_min, summary, desc):
        d = date.replace('-', '')
        t = lambda m: f'{m//60:02d}{m%60:02d}00'
        lines.extend(['BEGIN:VEVENT', f'UID:{uid}', f'DTSTAMP:{now}',
                      f'DTSTART;TZID=America/New_York:{d}T{t(start_min)}',
                      f'DTEND;TZID=America/New_York:{d}T{t(end_min)}',
                      f'SUMMARY:{ics_escape(summary)}',
                      f'LOCATION:{ics_escape(CAL_LOCATION)}',
                      f'DESCRIPTION:{ics_escape(desc)}',
                      'END:VEVENT'])
    for p in programs:
        ev(f'sacrament-{p["date"]}@ward-program', p['date'], 9*60, 10*60+30,
           'Milton Ward Sacrament Meeting',
           'Sacrament meeting. Full program: https://pouchman27.github.io/ward-program/')
    if act and act.get('rows'):
        for r in act['rows']:
            st = parse_time_min(r.get('schedule'))
            start = st if st is not None else 19*60
            bits = [b for b in [r.get('schedule'), r.get('notes'), r.get('reference')] if b]
            grp = '; '.join(f'{YW_CLASS_NAMES.get(k, k)}: {v}' for k, v in (r.get('groups') or {}).items())
            desc = '\n'.join(bits + ([grp] if grp else []))
            ev(f'activity-{r["date"]}@ward-program', r['date'], start, start + 60,
               'Milton Ward Youth Activity', desc)
    out = []
    for ln in lines:
        out.extend(ics_fold(ln))
    out.append('END:VCALENDAR')
    return '\r\n'.join(out) + '\r\n'

def cal_event_manifest(programs, act):
    events = []
    for p in programs:
        events.append({'uid': f'sacrament-{p["date"]}', 'date': p['date'], 'start_min': 9*60, 'end_min': 10*60+30,
                       'summary': 'Milton Ward Sacrament Meeting',
                       'description': 'Sacrament meeting. Full program: https://pouchman27.github.io/ward-program/'})
    if act and act.get('rows'):
        for r in act['rows']:
            st = parse_time_min(r.get('schedule'))
            start = st if st is not None else 19*60
            bits = [b for b in [r.get('schedule'), r.get('notes'), r.get('reference')] if b]
            grp = '; '.join(f'{YW_CLASS_NAMES.get(k, k)}: {v}' for k, v in (r.get('groups') or {}).items())
            events.append({'uid': f'activity-{r["date"]}', 'date': r['date'], 'start_min': start, 'end_min': start+60,
                           'summary': 'Milton Ward Youth Activity',
                           'description': '\n'.join(bits + ([grp] if grp else []))})
    return events

def write_calendar(programs, act, data):
    base = os.path.dirname(os.path.abspath(__file__))
    token_path = os.path.join(base, '.cal-token')
    token = os.environ.get('WARD_CAL_TOKEN', '').strip()
    if not token:
        if os.path.exists(token_path):
            token = open(token_path).read().strip()
        else:
            token = os.urandom(32).hex()
            with open(token_path, 'w') as f:
                f.write(token)
            os.chmod(token_path, 0o600)
    cal_name = f'cal-{token}.ics'
    data['calendar_url'] = f'https://pouchman27.github.io/ward-program/{cal_name}'
    site_dir = base
    for old in glob.glob(os.path.join(site_dir, 'cal-*.ics')):
        if os.path.basename(old) != cal_name:
            os.remove(old)
    with open(os.path.join(site_dir, cal_name), 'w', newline='') as f:
        f.write(build_ics(programs, act))
    return cal_name

def main():
    wb = openpyxl.load_workbook(XLSX)
    programs = []
    for name in wb.sheetnames:
        if name in SKIP_TABS: continue
        try:
            p = parse_program(wb[name])
            if p: programs.append(p)
        except Exception as e:
            print(f'WARN: failed tab {name}: {e}', file=sys.stderr)
    programs.sort(key=lambda p: p['date'], reverse=True)
    adv = {}
    if '2026 Advancements' in wb.sheetnames:
        try: adv = parse_advancements(Grid(wb['2026 Advancements']))
        except Exception as e: print(f'WARN advancements: {e}', file=sys.stderr)
    act = None
    act_path = '/tmp/activities.xlsx'
    if os.path.exists(act_path):
        try: act = parse_activities(act_path)
        except Exception as e: print(f'WARN activities: {e}', file=sys.stderr)
    data = {
        'generated_at': datetime.datetime.now(datetime.timezone.utc).isoformat(),
        'ward_name': 'Milton Ward',
        'sheet_url': SHEET_URL,
        'programs': programs,
        'advancements': adv,
        'activities': act,
    }
    ann_path = '/tmp/announcement_sections.json'
    if os.path.exists(ann_path):
        try: data['announcement_sections'] = json.load(open(ann_path))
        except Exception as e: print(f'WARN announcement sections: {e}', file=sys.stderr)
    cal_name = write_calendar(programs, act, data)
    with open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'cal_events.json'), 'w') as f:
        json.dump(cal_event_manifest(programs, act), f, indent=1)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    plaintext = json.dumps(data, indent=1).encode()
    # passphrase-gated payload: site/data.enc = salt16 + iv16 + AES-256-CBC(PBKDF2-HMAC-SHA256, 100k)
    from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
    from cryptography.hazmat.primitives import hashes, padding
    from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
    salt, iv = os.urandom(16), os.urandom(16)
    kdf = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=salt, iterations=100000)
    key = kdf.derive(SITE_PASSPHRASE.encode())
    padder = padding.PKCS7(128).padder()
    padded = padder.update(plaintext) + padder.finalize()
    ct = Cipher(algorithms.AES(key), modes.CBC(iv)).encryptor()
    enc = ct.update(padded) + ct.finalize()
    enc_path = os.path.join(os.path.dirname(OUT), 'data.enc')
    with open(enc_path, 'wb') as f:
        f.write(salt + iv + enc)
    if os.path.exists(OUT):  # never publish plaintext
        os.remove(OUT)
    print(f'OK: {len(programs)} programs, encrypted payload at {enc_path}, calendar at site/{cal_name}')

if __name__ == '__main__':
    main()
