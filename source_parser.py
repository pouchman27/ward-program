import openpyxl, json, re, datetime, os
HYMN_URLS=json.load(open(os.path.join(os.path.dirname(__file__),"hymn_map.json")))
SKIP_TABS={"Template","Sheet20","2026 Advancements"}
NUM_RE=re.compile(r"^\s*(\d+)\s*:\s*$")
def s(v):
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): v = int(v)
    if isinstance(v, (datetime.datetime, datetime.date)): return v.strftime('%-m/%-d/%Y')
    return str(v).strip()

QUOTEY=("we propose", "manifest it", "stand as your name", "bishop's office", "released the following", "have been called", "vote of thanks")

def is_noise(t):
    tl = t.lower().strip('"').strip()
    return not tl or any(q in tl for q in QUOTEY) or tl.startswith('"')

class Grid:
    def __init__(self, ws):
        self.ws = ws
        self.cells = {}   # (row,col) -> text
        self.links = {}   # (row,col) -> hyperlink target
        for row in ws.iter_rows():
            for c in row:
                t = s(c.value)
                if t:
                    self.cells[(c.row, c.column)] = t
                    if c.hyperlink and c.hyperlink.target:
                        self.links[(c.row, c.column)] = c.hyperlink.target
        self.max_row = max((r for r, _ in self.cells), default=0)

    def find(self, pattern, flags=re.I):
        rx = re.compile(pattern, flags)
        for (r, c), t in sorted(self.cells.items()):
            if rx.search(t):
                return r, c, t
        return None

    def right_of(self, row, col, max_col=14):
        for cc in range(col + 1, max_col + 1):
            t = self.cells.get((row, cc))
            if t: return t
        return ''

    def region_text(self, r1, r2, drop_markers=True, skip_cells=()):
        """Non-empty lines in row range, reading order, skipping noise/markers."""
        out = []
        for r in range(r1, r2 + 1):
            parts = []
            for c in range(1, 15):
                if (r, c) in skip_cells: continue
                t = self.cells.get((r, c), '')
                if not t or len(t) < 2: continue
                if drop_markers and NUM_RE.match(t): continue
                if is_noise(t): continue
                parts.append(t)
            if parts: out.append(' '.join(parts))
        return out

def parse_hymn(g, row):
    num, title = '', ''
    for c in range(1, 14):
        t = g.cells.get((row, c), '')
        if t.strip() == '#':
            v = g.right_of(row, c)
            if not re.match(r'^name\s*:', v, re.I):
                num = v
        elif re.match(r'^name\s*:', t, re.I):
            title = g.right_of(row, c)
    num = re.sub(r'\.0$', '', num)
    if not (num or title): return None
    out = {'number': num, 'title': title}
    if num:
        url = HYMN_URLS.get(str(num))
        if not url:
            raise ValueError(f'No official Church link mapped for hymn #{num}: {title}')
        out['url'] = url
    return out


def parse_program(ws):
    g = Grid(ws)
    d = g.find(r'^\d{4}-\d{2}-\d{2}|^\d{1,2}/\d{1,2}/\d{4}')
    date = None
    for (r, c), t in g.cells.items():
        if r <= 2:
            raw = ws.cell(row=r, column=c).value
            if isinstance(raw, (datetime.datetime, datetime.date)):
                date = raw.date() if isinstance(raw, datetime.datetime) else raw
                break
    if not date: return None
    prog = {'date': date.isoformat(),
            'slug': date.isoformat(),
            'tab': ws.title}
    # bulletin link omitted per Nathan, 2026-09-06
    p = g.find(r'^presiding$')
    if p: prog['presiding'] = g.right_of(p[0], p[1])
    cd = g.find(r'^conducting')
    if cd: prog['conducting'] = g.right_of(cd[0], cd[1])
    oh = g.find(r'opening hymn')
    if oh: prog['opening_hymn'] = parse_hymn(g, oh[0])
    inv = g.find(r'^invocation')
    if inv: prog['invocation'] = g.right_of(inv[0], inv[1])
    og = g.find(r'thank the organist')
    if og: prog['organist'] = g.right_of(og[0], og[1])
    ch = g.find(r'and chorister')
    if ch: prog['chorister'] = g.right_of(ch[0], ch[1])
    sh = g.find(r'sacrament hymn')
    ad = g.find(r'administration of the sacrament')
    ih = g.find(r'intermediate hymn')
    mn = g.find(r'special musical number')
    cl = g.find(r'closing hymn')
    be = g.find(r'^benediction')
    if sh: prog['sacrament_hymn'] = parse_hymn(g, sh[0])
    speakers = []
    for n in (1, 2, 3, 4):
        sp = g.find(rf'^speaker {n}\s*:')
        if sp:
            v = g.right_of(sp[0], sp[1])
            if v: speakers.append({'n': n, 'name': v})
    if speakers: prog['speakers'] = speakers
    if ih: prog['intermediate_hymn'] = parse_hymn(g, ih[0])
    if mn:
        entry = {}
        for rr in range(mn[0], mn[0] + 4):
            for cc in range(1, 14):
                t = g.cells.get((rr, cc), '')
                if re.match(r'^song name', t, re.I): entry['song'] = g.right_of(rr, cc)
                elif re.match(r'^singer name', t, re.I): entry['singers'] = g.right_of(rr, cc)
                elif re.match(r'^accompan', t, re.I): entry['accompanied'] = g.right_of(rr, cc)
        entry = {k: v for k, v in entry.items() if v}
        if entry: prog['musical_number'] = entry

    # The agenda sheet is the source of truth for the order between the
    # sacrament and closing hymn. Keep each populated item at its sheet row
    # instead of grouping all speakers ahead of the intermediate music.
    ordered = []
    for n in (1, 2, 3, 4):
        sp = g.find(rf'^speaker {n}\s*:')
        if sp:
            v = g.right_of(sp[0], sp[1])
            if v: ordered.append((sp[0], {'type': 'speaker', 'n': n, 'name': v}))
    if ih and prog.get('intermediate_hymn'):
        ordered.append((ih[0], {'type': 'hymn', 'label': 'Intermediate Hymn',
                                'hymn': prog['intermediate_hymn']}))
    if mn and prog.get('musical_number'):
        ordered.append((mn[0], {'type': 'musical_number', **prog['musical_number']}))
    if ordered:
        prog['program_order'] = [item for _, item in sorted(ordered, key=lambda x: x[0])]
    if cl: prog['closing_hymn'] = parse_hymn(g, cl[0])
    if be: prog['benediction'] = g.right_of(be[0], be[1])
    return prog



MONTHS={'jan': 1, 'feb': 2, 'mar': 3, 'apr': 4, 'may': 5, 'jun': 6, 'jul': 7, 'aug': 8, 'sept': 9, 'oct': 10, 'nov': 11, 'dec': 12}
